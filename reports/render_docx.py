"""
render_docx.py — แปลง ReportData เป็นไฟล์ Word (.docx)

ใช้ python-docx + ตั้งฟอนต์ไทย (Leelawadee UI มีติดมากับ Windows ทุกเครื่อง)
ได้ไฟล์ที่แก้ไขต่อเองได้ ต่างจาก PDF ที่แก้ไม่ได้
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from reports.report_data import ReportData

THAI_FONT = "Leelawadee UI"
SEV_LABEL = {"high": "สูง", "medium": "กลาง", "low": "ต่ำ"}
RED = RGBColor(0xDC, 0x26, 0x26)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
GRAY = RGBColor(0x64, 0x74, 0x8B)


def _set_thai_font(doc: Document) -> None:
    """บังคับฟอนต์ไทยทั้งเอกสาร (ต้องตั้ง eastasia/cs ด้วย ไม่งั้นไทยเพี้ยน)"""
    style = doc.styles["Normal"]
    style.font.name = THAI_FONT
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rpr.set(qn(attr), THAI_FONT)


def _delta_text(d: int) -> str:
    if d > 0:
        return f"▲ +{d}"
    if d < 0:
        return f"▼ {d}"
    return "–"


def _add_table(doc: Document, headers: list[str], rows: list[list[str]],
               red_cols: set[int] | None = None) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
            if red_cols and i in red_cols:
                run.font.color.rgb = RED
    doc.add_paragraph()


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    p.space_before = Pt(14)
    p.space_after = Pt(4)


def build(data: ReportData, out_path: str) -> str:
    doc = Document()
    _set_thai_font(doc)

    # ── ปก ──
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("รายงาน Pain Point การท่องเที่ยวพิษณุโลก")
    tr.bold = True
    tr.font.size = Pt(20)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(f"ประจำเดือน {data.period_label}\n"
                     f"วิเคราะห์จากรีวิว Google Maps · สร้างเมื่อ {data.generated_at}")
    sr.font.size = Pt(11)
    sr.font.color.rgb = GRAY

    if not data.has_data:
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("ไม่มีข้อมูลรีวิวในเดือนนี้").font.color.rgb = GRAY
        doc.save(out_path)
        return out_path

    # ── สรุปภาพรวม ──
    _heading(doc, "สรุปภาพรวม")
    _add_table(doc,
               ["ตัวชี้วัด", "จำนวน", "หมายเหตุ"],
               [["รีวิวในเดือนนี้", f"{data.reviews_in_month:,}", f"จาก {data.total_places} สถานที่"],
                ["คำบ่น (Pain Point)", f"{data.complaints_in_month:,}",
                 f"เดือนก่อน {data.prev_complaints:,} ({_delta_text(data.complaint_delta)})"],
                ["คำชม", f"{data.praise_in_month:,}", "จุดแข็งที่ควรรักษาไว้"],
                ["ปัญหาระดับสูง", f"{data.severity.get('high', 0):,}", "ต้องแก้เร่งด่วน"]])

    # ── 1. ปัญหาที่พบมากสุด ──
    _heading(doc, "1. ปัญหาที่พบมากที่สุด")
    _add_table(doc,
               ["หมวดปัญหา", "จำนวน", "สัดส่วน", "ระดับสูง", "เทียบเดือนก่อน"],
               [[f"{i+1}. {t['category']}", t["count"], f"{t['pct']}%", t["high"],
                 _delta_text(t["delta"])] for i, t in enumerate(data.top_problems)],
               red_cols={3})

    # ── 2. ระดับความรุนแรง ──
    sev_total = sum(data.severity.values()) or 1
    _heading(doc, "2. ระดับความรุนแรงของปัญหา")
    _add_table(doc, ["ระดับ", "จำนวน", "สัดส่วน"],
               [[SEV_LABEL.get(k, k), f"{v:,}", f"{v / sev_total * 100:.1f}%"]
                for k, v in sorted(data.severity.items(),
                                   key=lambda x: {"high": 0, "medium": 1, "low": 2}.get(x[0], 9))])

    # ── 3. โซน ──
    _heading(doc, "3. เปรียบเทียบตามพื้นที่")
    _add_table(doc, ["โซน", "คำบ่น", "คำชม", "ปัญหาหลัก"],
               [[z["zone"], z["complaints"], z["praise"], z["top_category"]] for z in data.zones])

    # ── 4. ร้านที่ถูกร้องเรียน ──
    _heading(doc, "4. สถานที่ที่ถูกร้องเรียนมากที่สุด")
    _add_table(doc, ["สถานที่", "โซน", "คำบ่น", "ระดับสูง", "ปัญหาหลัก"],
               [[f"{i+1}. {w['name']}", w["zone"], w["complaints"], w["high"], w["top_category"]]
                for i, w in enumerate(data.worst_places)],
               red_cols={3})

    # ── 5. จุดเด่น ──
    _heading(doc, "5. จุดเด่นที่ได้รับคำชม")
    _add_table(doc, ["หมวด", "คำชม"],
               [[f"{i+1}. {h['category']}", h["count"]] for i, h in enumerate(data.highlights)])

    # ── 6. ตัวอย่างรีวิว ──
    if data.sample_reviews:
        _heading(doc, f"6. ตัวอย่างรีวิว — {data.top_problems[0]['category']}")
        for s in data.sample_reviews:
            meta = doc.add_paragraph()
            mr = meta.add_run(f"{s['place']} · {'★' * (s['rating'] or 0)}")
            mr.font.size = Pt(9)
            mr.font.color.rgb = GRAY
            meta.space_after = Pt(0)
            body = doc.add_paragraph()
            body.paragraph_format.left_indent = Pt(18)
            br = body.add_run(s["text"])
            br.font.size = Pt(10)
            br.italic = True

    # ── ท้ายเอกสาร ──
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run("ระบบวิเคราะห์ Pain Point การท่องเที่ยวจังหวัดพิษณุโลก · มหาวิทยาลัยนเรศวร\n"
                      "ข้อมูลจากรีวิวสาธารณะบน Google Maps · วันที่รีวิวเป็นค่าโดยประมาณ")
    fr.font.size = Pt(9)
    fr.font.color.rgb = GRAY

    doc.save(out_path)
    return out_path
